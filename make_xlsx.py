import csv, datetime, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule, FormulaRule

GREEN = "21A366"
LIGHT = "EAF6EF"

def parse(v):
    if v is None or v == "":
        return v
    if v.startswith("="):
        return v
    try:
        if "." not in v and "," not in v and v.lstrip("-").isdigit():
            return int(v)
    except Exception:
        pass
    return v

def read_rows(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        text = f.read()
    first = next((l for l in text.splitlines() if l.strip()), "")
    sc, cm = first.find(";"), first.find(",")
    delim = ";" if (sc != -1 and (cm == -1 or sc < cm)) else ","
    if delim != ";" and sc == -1 and cm == -1:
        delim = ","
    import io
    return list(csv.reader(io.StringIO(text), delimiter=delim))

def build(csv_path, xlsx_path, sheet_name, date_cols=(), int_cols=(), cur_cols=(), pct_cols=(), cf=None):
    rows = read_rows(csv_path)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ri, row in enumerate(rows, 1):
        for ci, val in enumerate(row, 1):
            col = get_column_letter(ci)
            if ri > 1:
                if ci in date_cols:
                    try:
                        d = datetime.datetime.strptime(val, "%Y-%m-%d")
                        val = d
                    except Exception:
                        pass
                else:
                    val = parse(val)
            cell = ws.cell(row=ri, column=ci, value=val)
            cell.border = border
            if ri == 1:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=GREEN)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                if ri % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor=LIGHT)
                if ci in date_cols:
                    cell.number_format = "dd-mmm-yyyy"
                elif ci in cur_cols:
                    cell.number_format = '#,##0;[Red]-#,##0'
                elif ci in pct_cols:
                    cell.number_format = '0%'
                elif ci in int_cols:
                    cell.number_format = '0'

    # header + layout
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 22

    # column widths
    for ci in range(1, len(rows[0]) + 1):
        col = get_column_letter(ci)
        longest = max((len(str(rows[ri][ci - 1])) for ri in range(len(rows)) if ci - 1 < len(rows[ri])), default=8)
        ws.column_dimensions[col].width = min(max(longest + 4, 10), 42)

    # conditional formatting
    if cf:
        for rng, contains, color, fontcolor in cf:
            ws.conditional_formatting.add(
                rng,
                FormulaRule(formula=[f'ISNUMBER(SEARCH("{contains}",{rng.split(":")[0]}))'],
                            fill=PatternFill("solid", fgColor=color),
                            font=Font(color=fontcolor, bold=True))
            )

    wb.save(xlsx_path)
    print("ok", xlsx_path)

build("data/penjualan.csv",     "data/penjualan.xlsx",      "Penjualan",    date_cols=(1,), cur_cols=(5,6), int_cols=(4,))
build("data/gudang.csv",        "data/gudang.xlsx",         "Gudang",       int_cols=(3,4,5,6), cf=[("F2:F12","RESTOCK","FFC7CE","9C0006")])
build("data/nilai-siswa.csv",   "data/nilai-siswa.xlsx",    "Nilai",        int_cols=(2,), cf=[("D2:D12","REMED","FFC7CE","9C0006"),("D2:D12","LULUS","C6EFCE","006100")])
build("data/absensi.csv",       "data/absensi.xlsx",        "Absensi",      date_cols=(2,), int_cols=(4,), cf=[("E2:E12","SP","FFC7CE","9C0006")])
build("data/keuangan-warung.csv","data/keuangan-warung.xlsx","Keuangan",    cur_cols=(2,3,4,5), pct_cols=(6,))
build("data/pelanggan.csv",     "data/pelanggan.xlsx",       "Pelanggan",    cur_cols=(5,), cf=[("F2:F12","VIP","C6EFCE","006100")])
build("data/proyek.csv",        "data/proyek.xlsx",          "Proyek",       date_cols=(3,), int_cols=(5,6), cf=[("D2:D11","Selesai","C6EFCE","006100"),("D2:D11","Belum","FFC7CE","9C0006")])
